import io
import time

import pytest
from conftest import events, upload
from docx import Document as WordDocument
from sqlalchemy import select

from app import main, rag
from app.config import Settings, settings
from app.db import Chunk, Conversation, Document, SessionLocal
from app.rag import extract_text


def test_auth_cookie_and_cross_origin(client):
    assert client.get("/api/documents").status_code == 401
    response = client.post("/api/auth/demo", json={"account": "admin"})
    assert "HttpOnly" in response.headers["set-cookie"]
    assert "SameSite=lax" in response.headers["set-cookie"]
    assert client.post("/api/auth/logout", headers={"Origin": "https://evil.example"}).status_code == 403
    assert client.get("/api/auth/me").status_code == 200
    assert client.post("/api/auth/logout").status_code == 200
    assert client.get("/api/auth/me").status_code == 401


def test_upload_stream_citations_and_saved_history(admin):
    document = upload(admin).json()
    response = admin.post(
        "/api/chat", json={"question": "紫杉平台的测量记录保留多久？", "product": "紫杉平台", "version": "v1"}
    )
    assert response.headers["content-type"].startswith("text/event-stream")
    assert response.headers["x-accel-buffering"] == "no"
    stream = events(response)
    assert stream[0][0] == "conversation" and stream[-1][0] == "done"
    sources = next(data["items"] for kind, data in stream if kind == "sources")
    assert sources[0]["document_id"] == document["id"]
    answer = "".join(data["text"] for kind, data in stream if kind == "delta")
    assert "九十天" in answer and "[1]" in answer
    assert sum(kind == "delta" for kind, _ in stream) > 1
    conversation_id = stream[0][1]["id"]
    saved = admin.get(f"/api/conversations/{conversation_id}").json()["messages"]
    assert [m["role"] for m in saved] == ["user", "assistant"]
    assert saved[-1]["status"] == "complete" and saved[-1]["content"] == answer


def test_departments_filter_before_retrieval_and_sources(admin):
    doc = upload(admin, "猎鹰计划的内部口令为极光八号，只能向产品部展示。", visibility="department").json()
    private = upload(
        admin, "私有密钥说明仅用于管理员阅读，不应出现在普通成员的任何回答中。", visibility="private"
    ).json()
    admin.post("/api/auth/demo", json={"account": "support"})
    assert admin.get("/api/documents").json()["total"] == 0
    assert admin.get(f"/api/documents/{doc['id']}").status_code == 404
    stream = events(admin.post("/api/chat", json={"question": "猎鹰计划的口令是什么？"}))
    assert next(d["items"] for k, d in stream if k == "sources") == []
    assert "极光八号" not in str(stream)
    admin.post("/api/auth/demo", json={"account": "product"})
    assert admin.get("/api/documents").json()["total"] == 1
    assert admin.get(f"/api/documents/{private['id']}").status_code == 404
    assert admin.get(f"/api/documents/{doc['id']}").status_code == 200
    assert upload(admin).status_code == 403
    assert admin.delete(f"/api/documents/{doc['id']}").status_code == 403


def test_conversations_are_owned_by_current_user(admin):
    stream = events(admin.post("/api/chat", json={"question": "没有资料时如何回答？"}))
    conversation_id = stream[0][1]["id"]
    admin.post("/api/auth/demo", json={"account": "support"})
    assert admin.get("/api/conversations").json() == []
    assert admin.get(f"/api/conversations/{conversation_id}").status_code == 404
    assert admin.delete(f"/api/conversations/{conversation_id}").status_code == 404
    assert (
        admin.post("/api/chat", json={"conversation_id": conversation_id, "question": "继续"}).status_code
        == 404
    )


def test_product_and_version_filter_are_strict(admin):
    upload(admin)
    stream = events(admin.post("/api/chat", json={"question": "测量记录保留多久？", "product": "其他产品"}))
    assert next(d["items"] for k, d in stream if k == "sources") == []
    stream = events(
        admin.post(
            "/api/chat", json={"question": "测量记录保留多久？", "product": "紫杉平台", "version": "v9"}
        )
    )
    assert next(d["items"] for k, d in stream if k == "sources") == []


@pytest.mark.parametrize("change", ["delete", "replace", "revoke"])
def test_old_answers_redacted_after_source_changes(admin, change):
    doc_id = upload(admin).json()["id"]
    admin.post("/api/auth/demo", json={"account": "product"})
    stream = events(admin.post("/api/chat", json={"question": "测量记录保留多久？"}))
    conversation_id = stream[0][1]["id"]
    with SessionLocal() as db:
        doc = db.get(Document, doc_id)
        if change == "delete":
            db.delete(doc)
        elif change == "replace":
            doc.digest = "new-revision"
        else:
            doc.visibility = "private"
        db.commit()
    answer = admin.get(f"/api/conversations/{conversation_id}").json()["messages"][-1]
    assert answer["status"] == "stale"
    assert answer["sources"] == [] and "九十天" not in answer["content"]


def test_failed_replacement_keeps_old_index(admin, monkeypatch):
    doc = upload(admin).json()

    async def unavailable(*_args):
        raise RuntimeError("provider credentials must never escape")

    monkeypatch.setattr(main, "prepare_document", unavailable)
    response = admin.put(
        f"/api/documents/{doc['id']}",
        data={"title": "新标题", "product": "紫杉平台", "version": "v1"},
        files={"file": ("new.md", b"replacement text long enough")},
    )
    assert response.status_code == 502 and "credentials" not in response.text
    with SessionLocal() as db:
        assert db.get(Document, doc["id"]).title == doc["title"]
        assert db.scalar(select(Chunk).where(Chunk.document_id == doc["id"])) is not None


def test_duplicate_and_delete_cascade(admin):
    doc = upload(admin).json()
    assert upload(admin).status_code == 409
    assert admin.delete(f"/api/documents/{doc['id']}").status_code == 200
    with SessionLocal() as db:
        assert not db.scalars(select(Chunk).where(Chunk.document_id == doc["id"])).all()


def test_model_failure_emits_error_without_done_and_preserves_partial(admin, monkeypatch):
    upload(admin)

    async def broken(*_args):
        yield "部分回答"
        raise RuntimeError("secret-api-key")

    monkeypatch.setattr(main, "answer_stream", broken)
    response = admin.post("/api/chat", json={"question": "测量记录保留多久？"})
    stream = events(response)
    assert stream[-1][0] == "error" and not any(k == "done" for k, _ in stream)
    assert "secret-api-key" not in response.text
    saved = admin.get(f"/api/conversations/{stream[0][1]['id']}").json()["messages"][-1]
    assert saved["content"] == "部分回答" and saved["status"] == "error"
    with SessionLocal() as db:
        assert db.get(Conversation, stream[0][1]["id"]).busy_until == 0


def test_generation_lease_prevents_concurrent_answers(admin):
    stream = events(admin.post("/api/chat", json={"question": "测试"}))
    conversation_id = stream[0][1]["id"]
    with SessionLocal() as db:
        db.get(Conversation, conversation_id).busy_until = time.time() + 120
        db.commit()
    assert (
        admin.post("/api/chat", json={"question": "继续", "conversation_id": conversation_id}).status_code
        == 409
    )
    assert admin.delete(f"/api/conversations/{conversation_id}").status_code == 409


def test_user_creation_login_and_role_cannot_be_escalated(admin):
    response = admin.post(
        "/api/users",
        json={
            "username": "new-user",
            "name": "新同事",
            "department": "市场部",
            "password": "a-long-password-123",
            "role": "admin",
        },
    )
    assert response.status_code == 201 and response.json()["role"] == "member"
    assert "password" not in response.text
    admin.post("/api/auth/logout")
    assert (
        admin.post("/api/auth/login", json={"username": "new-user", "password": "wrong"}).status_code == 401
    )
    assert (
        admin.post(
            "/api/auth/login", json={"username": "new-user", "password": "a-long-password-123"}
        ).status_code
        == 200
    )
    assert admin.get("/api/users").status_code == 403


def test_docx_parser_includes_tables_and_rejects_empty():
    doc = WordDocument()
    doc.add_paragraph("产品使用说明，请按以下表格配置")
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text, row.cells[1].text = "数据保留", "九十天"
    buffer = io.BytesIO()
    doc.save(buffer)
    assert "九十天" in extract_text("guide.docx", buffer.getvalue())
    with pytest.raises(ValueError, match="足够文字"):
        extract_text("empty.md", b"abc")
    with pytest.raises(ValueError, match="仅支持"):
        extract_text("bad.exe", b"executable")


def test_real_mode_rejects_insecure_configuration():
    with pytest.raises(ValueError, match="SESSION_SECRET"):
        Settings(_env_file=None, demo_mode=False, session_secret="local-development")


async def test_upstream_truncated_stream_is_not_treated_as_success(monkeypatch):
    import httpx

    original = httpx.AsyncClient
    monkeypatch.setattr(settings, "demo_mode", False)
    monkeypatch.setattr(settings, "model_base_url", "https://model.invalid/v1")

    def upstream(_request):
        return httpx.Response(
            200,
            text='data: {"choices":[{"delta":{"content":"partial"}}]}\n\n',
            headers={"Content-Type": "text/event-stream"},
        )

    monkeypatch.setattr(
        rag.httpx, "AsyncClient", lambda **_kwargs: original(transport=httpx.MockTransport(upstream))
    )
    chunks = []
    with pytest.raises(ValueError, match="提前结束"):
        async for chunk in rag.answer_stream(
            "问题", [{"number": 1, "title": "资料", "product": "示例", "version": "v1", "text": "依据"}], []
        ):
            chunks.append(chunk)
    assert chunks == ["partial"]


async def test_embedding_dimension_validation(monkeypatch):
    import httpx

    original = httpx.AsyncClient
    monkeypatch.setattr(settings, "demo_mode", False)
    monkeypatch.setattr(settings, "model_base_url", "https://model.invalid/v1")
    monkeypatch.setattr(
        rag.httpx,
        "AsyncClient",
        lambda **_kwargs: original(
            transport=httpx.MockTransport(
                lambda _r: httpx.Response(200, json={"data": [{"index": 0, "embedding": [1, 2]}]})
            )
        ),
    )
    with pytest.raises(ValueError, match="维度"):
        await rag.embed(["text"])


def test_relative_sqlite_path_is_independent_of_working_directory():
    from app.config import ROOT

    config = Settings(_env_file=None, database_url="sqlite:///./data/portable.db", demo_mode=True)
    assert config.database_url == f"sqlite:///{ROOT / 'data' / 'portable.db'}"


def test_demo_seed_marker_is_stored_in_database(admin):
    from app.db import AppState
    from app.seed import seed

    seed()
    with SessionLocal() as db:
        assert db.get(AppState, "demo-seeded") is not None
        sample = db.scalar(select(Document).where(Document.filename == "getting-started.md"))
        sample_id = sample.id
        db.delete(sample)
        db.commit()
    seed()
    with SessionLocal() as db:
        assert db.get(Document, sample_id) is None
        assert db.scalar(select(Document).where(Document.filename == "getting-started.md")) is None


async def test_semantic_retrieval_respects_product_scope(admin, monkeypatch):
    from app.db import User

    admin_id = admin.get("/api/auth/me").json()["id"]
    doc = upload(admin, "这份文档讨论柚木系统的独特功能与操作流程。", product="柚木系统").json()
    other = upload(admin, "这份文档讨论桂花系统的内部功能与操作流程。", product="桂花系统").json()
    monkeypatch.setattr(settings, "demo_mode", False)

    async def embeddings(_texts):
        return [[0.0, 1.0, 0.0]]

    monkeypatch.setattr(rag, "embed", embeddings)
    with SessionLocal() as db:
        for doc_id in [doc["id"], other["id"]]:
            db.get(Document, doc_id).embedding_profile = rag.embedding_profile()
            for chunk in db.scalars(select(Chunk).where(Chunk.document_id == doc_id)):
                chunk.embedding = [0.0, 1.0, 0.0]
        db.commit()
        # No lexical overlap: this result must come from the semantic branch.
        results = await rag.retrieve(db, db.get(User, admin_id), "数字化升级", "柚木系统", "")
        assert len(results) == 1 and results[0]["document_id"] == doc["id"]
